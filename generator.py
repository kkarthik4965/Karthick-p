"""Formatting utilities: sanitize, HTML preview, DOCX and PDF export."""
import html
import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from fpdf import FPDF

import config

_REPLACEMENTS = {
    "\u2018": "'", "\u2019": "'", "\u201c": '"', "\u201d": '"',
    "\u2013": "-", "\u2014": "-", "\u2026": "...", "\u2022": "-",
    "\u00a0": " ", "\u2192": "->", "\u2713": "", "\u2714": "",
}


def sanitize_text(text: str) -> str:
    """Remove typographic characters and markdown symbols for clean output."""
    for old, new in _REPLACEMENTS.items():
        text = text.replace(old, new)
    text = text.replace("**", "").replace("`", "")
    text = re.sub(r"^\s*#{1,6}\s*", "", text, flags=re.MULTILINE)
    text = re.sub(r"^\s*\*\s+", "- ", text, flags=re.MULTILINE)
    return text.strip()


def _split_terms(terms: str) -> list[str]:
    return [t.strip() for t in (terms or "").split(";") if t.strip()]


def _is_heading(line: str) -> bool:
    s = line.strip()
    if not s or len(s) > 80 or s.startswith(("-", "(")) or re.match(r"^[a-z]\)", s):
        return False
    return s.endswith(":") or (s.isupper() and len(s) > 3)


def _is_bullet(line: str) -> bool:
    return line.strip().startswith("- ")


def format_html_preview(text: str) -> str:
    parts, in_list = [], False
    for raw in sanitize_text(text).splitlines():
        line = raw.strip()
        if in_list and not _is_bullet(line):
            parts.append("</ul>")
            in_list = False
        if not line:
            continue
        safe = html.escape(line)
        if _is_bullet(line):
            if not in_list:
                parts.append("<ul>")
                in_list = True
            parts.append(f"<li>{html.escape(line[2:].strip())}</li>")
        elif _is_heading(line):
            parts.append(f"<h4 style='margin:14px 0 4px 0;color:#ffffff'>{safe}</h4>")
        else:
            parts.append(f"<p style='margin:4px 0;line-height:1.5'>{safe}</p>")
    if in_list:
        parts.append("</ul>")
    return "".join(parts)


def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    text = sanitize_text(text)
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    logo = doc.add_paragraph()
    logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        logo.add_run().add_picture(config.LOGO_PATH, width=Inches(2.2))
    except Exception:
        pass

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(sanitize_text(doc_type).upper())
    run.bold = True
    run.font.size = Pt(16)

    items = _split_terms(terms)
    if items:
        p = doc.add_paragraph()
        p.add_run("Summary of Key Terms").bold = True
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "No."
        table.rows[0].cells[1].text = "Term"
        for i, t in enumerate(items, 1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = sanitize_text(t)
        doc.add_paragraph()

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        if _is_bullet(line):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(line)
            r.bold = _is_heading(line)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run(config.FOOTER_TEXT)
    fr.italic = True
    fr.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _latin(s: str) -> str:
    return sanitize_text(s).encode("latin-1", "replace").decode("latin-1")


class _PDF(FPDF):
    def __init__(self, doc_type: str):
        super().__init__()
        self.doc_type = _latin(doc_type)

    def header(self):
        try:
            self.image(config.LOGO_PATH, x=(self.w - 40) / 2, y=8, w=40)
        except Exception:
            pass
        self.set_y(22)
        self.set_font("Helvetica", "B", 13)
        self.cell(0, 8, self.doc_type, align="C", new_x="LMARGIN", new_y="NEXT")
        self.ln(4)

    def footer(self):
        self.set_y(-15)
        self.set_font("Helvetica", "I", 8)
        self.cell(0, 10, config.FOOTER_TEXT, align="C")


def format_pdf(text: str, doc_type: str, terms: str = "") -> bytes:
    pdf = _PDF(doc_type)
    pdf.set_margins(20, 20, 20)
    pdf.set_auto_page_break(auto=True, margin=20)
    pdf.add_page()

    def write(line, style="", indent=0):
        pdf.set_font("Helvetica", style, 11)
        pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(0, 6, line, new_x="LMARGIN", new_y="NEXT")

    items = _split_terms(terms)
    if items:
        write("Summary of Key Terms", "B")
        for t in items:
            write("- " + _latin(t), indent=5)
        pdf.ln(3)

    for raw in sanitize_text(text).splitlines():
        line = _latin(raw).strip()
        if not line:
            pdf.ln(2)
        elif _is_bullet(line):
            write("- " + line[2:].strip(), indent=5)
        elif _is_heading(line):
            pdf.ln(1)
            write(line, "B")
        else:
            write(line)
    return bytes(pdf.output())
